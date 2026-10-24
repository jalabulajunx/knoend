#!/usr/bin/env python3
"""Builds the small PowerPoint deck the companion's tests read.

The deck is a made-up pitch for OpenPager (Kn°End's earlier name) with speaker notes: a short
note, a long one that has to be paginated for the board, and a slide with no notes. Nothing private.

    pip install python-pptx
    python3 tools/make_test_pptx.py   # writes companion/crates/cli/tests/fixtures/openpager-demo.pptx
"""

from pathlib import Path

from pptx import Presentation

OUT = Path(__file__).resolve().parent.parent / "companion/crates/cli/tests/fixtures/openpager-demo.pptx"

# (title, body, speaker notes); None means no notes on that slide.
SLIDES = [
    (
        "OpenPager",
        "A shortcut keypad and pager on a cheap ESP32-S3 touch board",
        "OpenPager intro: hold up the board, say what it is in one line, then move on.",
    ),
    (
        "The problem",
        "Keypads are closed, skip Linux, and every app needs a profile",
        "Ask who owns a macro pad.\nAsk who set it up for more than one app.",
    ),
    (
        "How it works",
        "Board + companion service + layouts as data",
        "The board works alone as a Bluetooth keyboard, with nothing installed on the computer.\n"
        "The companion service watches which app, website or terminal program is in front and "
        "pushes the matching keys to every connected board.\n"
        "Layouts are TOML files checked by one validator, whether built in, shared by someone else "
        "or drafted by an AI for an app nobody has set up yet.\n"
        "Over Wi-Fi it becomes a pager: approve a coding agent's command from another room.",
    ),
    ("Live demo", "Switch apps and watch the keys change", None),
    (
        "Thank you",
        "github.com/jalabulajunx/knoend",
        "Point to the repository. Take questions.",
    ),
]


def main() -> None:
    deck = Presentation()
    layout = deck.slide_layouts[1]  # title and content
    for title, body, notes in SLIDES:
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes is not None:
            slide.notes_slide.notes_text_frame.text = notes
    OUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
